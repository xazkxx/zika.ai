# -*- coding: utf-8 -*-
"""浩技云数据 → ERP 标准 Excel 转换
输入: haoji_data/raw_*.json
输出: D:/ERP软件/migration/erp_products|customers|suppliers|warehouses|opening_stock.xlsx
      D:/ERP软件/migration/备查_*.xlsx  （银行账户/员工/合同/报价/调拨/BOM/成品入库）
用 D 盘 venv 跑（需要 openpyxl）:
  D:/ERP软件/.venv/Scripts/python.exe haoji_to_erp.py
"""
import sys, io, os, json, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.path.insert(0, 'D:/ERP软件')
from openpyxl import Workbook
from openpyxl.styles import Font
from app.services.excel_io import COLUMNS

RAW = 'haoji_data'
OUT = 'D:/ERP软件/migration'
os.makedirs(OUT, exist_ok=True)

def load(name):
    return json.load(open(f'{RAW}/raw_{name}.json', encoding='utf-8'))

def fnum(v):
    if v in (None, ''): return None
    try: return float(str(v).strip())
    except ValueError: return None

def fint(v):
    f = fnum(v)
    return None if f is None else int(round(f))

def fstr(v):
    if v is None: return None
    s = str(v).strip()
    return s or None

def clean_addr(s):
    """'地址:xxx 邮政编码:710054' -> ('xxx', '710054')"""
    if not s: return None, None
    t = str(s).strip()
    post = None
    m = re.search(r'邮政编码[:：]?\s*(\d+)', t)
    if m:
        post = m.group(1)
        t = t[:m.start()] + t[m.end():]
    t = re.sub(r'^地址[:：]?', '', t).strip(' ;；,，')
    return (t or None), post

def write_import_excel(kind, records, path):
    """按 excel_io 列规格写导入 Excel：第 1 行中文表头，数据从第 2 行"""
    specs = COLUMNS[kind]
    wb = Workbook(); ws = wb.active; ws.title = kind
    for c, sp in enumerate(specs, 1):
        cell = ws.cell(row=1, column=c, value=sp['label'])
        cell.font = Font(bold=True)
    for r, rec in enumerate(records, 2):
        for c, sp in enumerate(specs, 1):
            v = rec.get(sp['field'])
            if v is not None and v != '':
                ws.cell(row=r, column=c, value=v)
    # 列宽
    for c, sp in enumerate(specs, 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = max(10, len(sp['label']) * 2.2)
    wb.save(path)
    print(f'  saved {path}  ({len(records)} 行)')

def write_backup_excel(name, data, path):
    """备查表：中文列头（动态字段映射，无映射用原名）"""
    cols = data['columns']; rows = data['rows']
    # 汇总所有键，保持首行出现顺序
    keys = []
    for r in rows:
        for k in r:
            if k not in keys: keys.append(k)
    skip = {'state', 'createName', 'createBy', 'sysOrgCode'}
    keys = [k for k in keys if k not in skip and k != 'id'] or keys
    wb = Workbook(); ws = wb.active; ws.title = name[:30]
    headers = [(cols.get(k, k) if k != 'id' else '浩技云ID') for k in keys]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
    for r, row in enumerate(rows, 2):
        for c, k in enumerate(keys, 1):
            v = row.get(k)
            if v not in (None, ''):
                if k == 'id': v = str(v)[-8:]
                ws.cell(row=r, column=c, value=v)
    for c in range(1, len(keys) + 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 16
    wb.save(path)
    print(f'  saved {path}  ({len(rows)} 行)')

# ---------------------------------------------------------------- 商品
def conv_products():
    data = load('products')
    groups = {}
    order = []
    for r in data['rows']:
        n = fstr(r.get('proNum'))
        if not n: continue
        if n not in groups:
            groups[n] = []; order.append(n)
        groups[n].append(r)
    def richness(r):
        return sum(1 for v in r.values() if v not in ('', None))

    META = {'id', 'createName', 'createBy', 'createDate', 'sysOrgCode', 'state', 'status'}

    def core(r):
        """严格签名：除元字段外的所有字段值（含空值）全部一致才算重复。
        用户要求：重复编号内只要有不一致的内容就保留。"""
        return tuple(sorted((k, str(v if v is not None else '').strip())
                            for k, v in r.items() if k not in META))
    records, sku_map, dup_log = [], {}, []
    for n in order:
        g = sorted(groups[n], key=richness, reverse=True)
        main = g[0]
        sku_map[n] = n  # 库存引用 -> 主条目
        records.append({
            'sku': n,
            'name': fstr(main.get('proZnName')),
            'category': fstr(main.get('proTypeName')),
            'unit': fstr(main.get('unit')),
            'standards': fstr(main.get('standards')),
            'brand': fstr(main.get('brand')),
            'mnemonic': fstr(main.get('proZjm')),
            'bar_code': fstr(main.get('barCode')),
            'color': fstr(main.get('a01a01a01')),
            'size': fstr(main.get('a01a01a02')),
            'manufacturer': fstr(main.get('a01a01a03')),
            'serial_no': fstr(main.get('a01a01a05')),
            'stock_upper': fint(main.get('a01a01a06')),
            'stock_lower': fint(main.get('a01a01a07')),
            'pack_qty': fint(main.get('a01a01a09')),
            'shelf_life': fint(main.get('a01a01a17')),
            'purchase_price': fnum(main.get('costPrice')),
            'sale_price': fnum(main.get('sellPrice')),
            'note': None,
        })
        idx = 2
        seen = {core(main)}
        for other in g[1:]:
            s = core(other)
            if s in seen:
                continue  # 与已保留档案完全一致 → 合并
            seen.add(s)
            sku = f'{n}-{idx}'; idx += 1
            records.append({
                'sku': sku,
                'name': fstr(other.get('proZnName')),
                'category': fstr(other.get('proTypeName')),
                'unit': fstr(other.get('unit')),
                'standards': fstr(other.get('standards')),
                'brand': fstr(other.get('brand')),
                'mnemonic': fstr(other.get('proZjm')),
                'bar_code': fstr(other.get('barCode')),
                'color': fstr(other.get('a01a01a01')),
                'size': fstr(other.get('a01a01a02')),
                'manufacturer': fstr(other.get('a01a01a03')),
                'purchase_price': fnum(other.get('costPrice')),
                'sale_price': fnum(other.get('sellPrice')),
                'pack_qty': fint(other.get('a01a01a09')),
                'note': f'浩技云重复编号 {n}，差异档案',
            })
            dup_log.append(f'{n} -> {sku}')
    write_import_excel('products', records, f'{OUT}/erp_products.xlsx')
    print('  重复编号处理:', dup_log or '无')
    return records

# ---------------------------------------------------------------- 客户
def conv_customers():
    data = load('customers')
    records = []
    for r in data['rows']:
        addr, post = clean_addr(r.get('address'))
        note = f'浩技云类型:{r.get("cusType")}' if r.get('cusType') else None
        if post: note = (note + '；' if note else '') + f'邮编:{post}'
        records.append({
            'code': fstr(r.get('cusNum')),
            'name': fstr(r.get('cusName')),
            'address': addr,
            'contact': fstr(r.get('zlxr')),
            'contact_phone': fstr(r.get('workphone')),
            'businesser': fstr(r.get('businesserName')),
            'tax_no': fstr(r.get('suiNum')),
            'bank_name': fstr(r.get('bankName')),
            'bank_account': fstr(r.get('bankAccount')),
            'level': fstr(r.get('cusLevel')),
            'discount': fnum(r.get('cusZk')),
            'credit_limit': fnum(r.get('xyed')),
            'note': note,
        })
    write_import_excel('customers', records, f'{OUT}/erp_customers.xlsx')

# ---------------------------------------------------------------- 供应商
def conv_suppliers():
    data = load('suppliers')
    records = []
    for r in data['rows']:
        records.append({
            'code': fstr(r.get('companyCode')),
            'name': fstr(r.get('companyNameZn')),
            'short_name': fstr(r.get('zjm')),
            'phone': fstr(r.get('primaryContactTel')),
            'email': fstr(r.get('primaryContactEmail')),
            'address': fstr(r.get('addressZn')),
            'contact': fstr(r.get('primaryContact')),
            'main_contact': fstr(r.get('primaryContact')),
            'mobile': fstr(r.get('primaryContactTel')),
            'bank_name': fstr(r.get('bankName')),
            'bank_account': fstr(r.get('bankAccount')),
            'init_payable': fnum(r.get('bcqkMoney')),
        })
    write_import_excel('suppliers', records, f'{OUT}/erp_suppliers.xlsx')

# ---------------------------------------------------------------- 仓库 + 期初库存
WAREHOUSE_CODE = {}
def conv_warehouses():
    data = load('warehouses')
    records = []
    for i, r in enumerate(data['rows'], 1):
        code = f'CK-{i:03d}'
        WAREHOUSE_CODE[fstr(r.get('storageName'))] = code
        records.append({
            'code': code,
            'name': fstr(r.get('storageName')),
            'location': fstr(r.get('storageAddress')),
            'note': fstr(r.get('remark')),
        })
    write_import_excel('warehouses', records, f'{OUT}/erp_warehouses.xlsx')

def conv_opening_stock(products):
    data = load('stock')
    by_sku = {p['sku']: p for p in products}
    records, skipped = [], []
    for r in data['rows']:
        num = fstr(r.get('a01a09a08'))
        wh = WAREHOUSE_CODE.get(fstr(r.get('storageName')))
        qty = fnum(r.get('a01a09a02'))
        price = fnum(r.get('a01a09a13'))
        if not num or not wh or qty is None:
            skipped.append(num); continue
        sku = num if num in by_sku else (num.split('-')[0] if num.split('-')[0] in by_sku else None)
        if not sku:
            skipped.append(num); continue
        records.append({
            'warehouse_code': wh,
            'sku': sku,
            'quantity': int(round(qty)),
            'unit_price': price,
            'note': '浩技云库存迁移 2026-10',
        })
    write_import_excel('opening-stock', records, f'{OUT}/erp_opening_stock.xlsx')
    if skipped: print('  期初库存跳过:', skipped)

# ---------------------------------------------------------------- 备查表
def conv_backups():
    for name in ['bank', 'staff', 'contract_cg', 'contract_xs', 'offers', 'allot', 'bom', 'scstorage']:
        data = load(name)
        write_backup_excel(name, data, f'{OUT}/备查_{name}.xlsx')

if __name__ == '__main__':
    print('== 转换商品')
    products = conv_products()
    print('== 转换客户'); conv_customers()
    print('== 转换供应商'); conv_suppliers()
    print('== 转换仓库'); conv_warehouses()
    print('== 转换期初库存')
    conv_opening_stock(products)
    print('== 备查表'); conv_backups()
    print('完成 ->', OUT)
