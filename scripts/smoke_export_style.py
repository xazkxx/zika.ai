# -*- coding: utf-8 -*-
"""导出精装冒烟：取库里真实对账单，走 PDF / Word / Excel 三条链路验证样式。

2026-10-26（v1.77.0「增强对账单 UI 重构」）：
  Excel 这一条**不再写死行号**。抬头多了"公司名"一行、下面还有汇总卡片条，
  表头位置会随（公司名 / 卡片）浮动 —— 旧脚本把表头钉在第 3~8 行，
  一旦抬头变化就误报"表头没找到"。现在一律**按第一列文字定位表头**。

期望（与网页端同一套视觉语言）：
  · 抬头：公司名（取自 /api/v1/settings/company 的 own_name）/ 单据名 / 共 N 条 + 导出时间；
  · 表头 #1F4E78 深蓝底 + 白字，冻结在表头下第一行；
  · 斑马纹 #F7F9FC，从第 2 条数据开始；
  · 列宽按内容自适应（8~46 之间，中文按 2 字宽算）；
  · 金额是**数值** + #,##0.00（千分位），绝不能被写成文本；
  · 打印：A4 横版 / 按一页宽缩放 / 表头重复 / 页脚页码 / 页眉打印时间域。

跑法：python scripts/smoke_export_style.py
"""

import base64
import io
import sys

from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)
r = client.post("/api/v1/auth/login", json={"username": "admin", "password": "Admin123"})
token = {"Authorization": f"Bearer {r.json()['access_token']}"}

# 找一张有明细的对账单
rows = client.get("/api/v1/statements?limit=50", headers=token).json()
rows = rows if isinstance(rows, list) else rows.get("items") or []
if not rows:
    print("NO_STATEMENT")
    sys.exit(0)
sid = rows[0]["id"]
print(f"statement id={sid} no={rows[0].get('statement_no')}")

problems: list[str] = []


def check(ok: bool, label: str, detail: str = "") -> bool:
    if ok:
        print(f"  OK   {label}{('  — ' + detail) if detail else ''}")
    else:
        print(f"  FAIL {label}{('  — ' + detail) if detail else ''}")
        problems.append(label)
    return ok


# ---- PDF：必须是 A4 横版 + 深蓝表头填充 ----
resp = client.get(f"/api/v1/print/statement-record/{sid}", headers={**token, "Accept": "application/json"})
assert resp.status_code == 200, resp.text
pdf = base64.b64decode(resp.json()["content_base64"])
import pymupdf as fitz

doc = fitz.open(stream=pdf, filetype="pdf")
page = doc[0]
w, h = page.rect.width, page.rect.height
print(f"PDF: {len(doc)} pages, page={w:.1f}x{h:.1f}")
check(w > h, "PDF 横版")
found_fill = False
for p in doc:
    for dr in p.get_drawings():
        f = dr.get("fill")
        if f and abs(f[0] - 31 / 255) < 0.02 and abs(f[1] - 78 / 255) < 0.02 and abs(f[2] - 120 / 255) < 0.02:
            found_fill = True
            break
    if found_fill:
        break
check(found_fill, "PDF 深蓝表头填充 #1F4E78")
doc.close()

# ---- Word：横版 section + 汇总卡片 ----
resp = client.get(f"/api/v1/exports/word/statement/{sid}", headers=token)
assert resp.status_code == 200, resp.text
from docx import Document as WDoc

wd = WDoc(io.BytesIO(resp.content))
sec = wd.sections[0]
print(f"WORD: orientation={sec.orientation}, page={sec.page_width.cm:.1f}x{sec.page_height.cm:.1f}cm")
check(sec.page_width.cm > sec.page_height.cm, "Word 横版")
all_head = "".join(c.text for t in wd.tables[:3] for c in t.rows[0].cells)
check("期初余额" in all_head, "Word 汇总卡片", all_head[:30])

# ---- Excel：抬头 / 深蓝表头 / 斑马纹 / 冻结 / 自适应列宽 / 千分位 / 打印设置 ----
resp = client.get("/api/v1/exports/statements", headers=token)
assert resp.status_code == 200, resp.text
from openpyxl import load_workbook

wb = load_workbook(io.BytesIO(resp.content))
check("对账单" in wb.sheetnames and "对账明细" in wb.sheetnames, "两张工作表", str(wb.sheetnames))
ws = wb["对账单"]


def find_row(label: str, limit: int = 14) -> int | None:
    """按第一列文字找行号 —— 抬头/卡片一变，行号就会变，不能写死。"""
    for row in range(1, limit + 1):
        if str(ws.cell(row=row, column=1).value or "") == label:
            return row
    return None


head_row = find_row("对账单号")
if not check(head_row is not None, "找到表头行"):
    print("SMOKE_FAIL")
    sys.exit(1)

# 抬头公司名（当前账套 own_name）
company = str((client.get("/api/v1/settings/company", headers=token).json() or {}).get("own_name") or "")
first = str(ws.cell(row=1, column=1).value or "")
print(f"EXCEL first cell: {first!r} / company={company!r}")
if company:
    check(first == company, "抬头公司名", first)
title_cell = str(ws.cell(row=2 if company else 1, column=1).value or "")
check("对账单台账" in title_cell, "单据名行", title_cell[:20])
check(find_row("期初余额") is not None, "汇总卡片条（期初余额…）")

# 深蓝表头 + 白字
head_cell = ws.cell(row=head_row, column=1)
fill = head_cell.fill.fgColor.rgb
print(f"EXCEL header at row {head_row}, fill={fill}")
check(str(fill)[-6:] == "1F4E78", "表头深蓝 #1F4E78", str(fill))
check(str(head_cell.font.color.rgb)[-6:] == "FFFFFF", "表头白字", str(head_cell.font.color.rgb))

# 冻结在表头下第一行
check(ws.freeze_panes == f"A{head_row + 1}", "冻结表头", str(ws.freeze_panes))

# 斑马纹：偶数位数据行有浅底，奇数位没有
first_data = ws.cell(row=head_row + 1, column=1)
second_data = ws.cell(row=head_row + 2, column=1)
zebra = second_data.fill.fgColor.rgb
print(f"EXCEL zebra fill row {head_row + 2}: {zebra}")
check(str(zebra)[-6:] == "F7F9FC", "斑马纹 #F7F9FC", str(zebra))
check(str(first_data.fill.fgColor.rgb)[-6:] != "F7F9FC", "首行不上斑马纹")

# 千分位 + 数值：金额列必须是数字，且格式带千分位
labels = [str(ws.cell(row=head_row, column=c).value or "") for c in range(1, ws.max_column + 1)]
if "期初余额" in labels:
    money_cell = ws.cell(row=head_row + 1, column=labels.index("期初余额") + 1)
    print(f"EXCEL money cell: {money_cell.value!r} fmt={money_cell.number_format}")
    check(isinstance(money_cell.value, (int, float)), "金额是数值（非文本）", repr(money_cell.value))
    check(money_cell.number_format == "#,##0.00", "金额千分位格式", money_cell.number_format)
# 列宽自适应：按内容重算一遍，和实际宽度对齐（±1）
def display_width(value) -> int:
    if value is None or value == "":
        return 0
    return sum(2 if ord(ch) > 127 else 1 for ch in str(value))


mismatch: list[str] = []
for idx, label in enumerate(labels, start=1):
    needed = max(display_width(label) + 2, 8)
    for row in range(head_row + 1, min(ws.max_row, head_row + 40) + 1):
        needed = max(needed, display_width(ws.cell(row=row, column=idx).value) + 2)
    needed = min(needed, 46)  # 长文本列上限，与 doc_export._MAX_COL_WIDTH 一致
    actual = ws.column_dimensions[ws.cell(row=head_row, column=idx).column_letter].width
    if actual is None or abs(float(actual) - needed) > 1.0:
        mismatch.append(f"{label}:{actual}≠{needed}")
print("EXCEL widths:", [round(float(ws.column_dimensions[ws.cell(row=head_row, column=c).column_letter].width or 0), 1) for c in range(1, min(len(labels), 8) + 1)])
check(not mismatch, "列宽按内容自适应", "；".join(mismatch[:4]))

# 打印设置
check(ws.page_setup.orientation == "landscape", "打印 A4 横版", str(ws.page_setup.orientation))
check(ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0, "按一页宽缩放")
check(getattr(ws.sheet_properties.pageSetUpPr, "fitToPage", False) is True, "fitToPage 打开")
check(
    (ws.print_title_rows or "").replace("$", "") == f"{head_row}:{head_row}",
    "每页重复表头",
    str(ws.print_title_rows),
)
check("&P" in (ws.oddFooter.center.text or ""), "页脚页码域", str(ws.oddFooter.center.text))
check("&D" in (ws.oddHeader.right.text or ""), "页眉打印时间域", str(ws.oddHeader.right.text))

print("SMOKE_OK" if not problems else f"SMOKE_FAIL {problems}")
sys.exit(0 if not problems else 1)
