"""生成「智能识别上传」验收用的样本文件（2026-10-23）。

验收脚本需要**真的图片/Word**才能证明 OCR 和 python-docx 都通，所以样本现场生成，
不往仓库里塞二进制。三份：

* `license.png`    —— 模拟营业执照（PIL + 微软雅黑，白底黑字）
* `namecard.png`   —— 模拟名片（联系人/手机/邮箱/地址）
* `quote.docx`     —— 模拟报价单（正文 + 商品明细表格）

跑法：`./.venv/Scripts/python.exe scripts/make_smart_upload_fixtures.py`
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "验收样本"

FONT = "C:/Windows/Fonts/msyh.ttc"


def main() -> int:
    from PIL import Image, ImageDraw, ImageFont

    if not Path(FONT).exists():
        print(f"找不到中文字体 {FONT}，无法生成图片样本", file=sys.stderr)
        return 1
    OUT.mkdir(parents=True, exist_ok=True)

    big = ImageFont.truetype(FONT, 44)
    mid = ImageFont.truetype(FONT, 30)

    # —— 营业执照 ——
    license_img = Image.new("RGB", (1240, 760), "white")
    draw = ImageDraw.Draw(license_img)
    draw.text((420, 30), "营 业 执 照", font=big, fill="black")
    draw.text((540, 96), "（副本）", font=ImageFont.truetype(FONT, 24), fill="black")
    lines = [
        "统一社会信用代码  91610103MA6TX8K23Q",
        "名        称  陕西送变电工程有限公司",
        "类        型  有限责任公司（非自然人投资或控股的法人独资）",
        "住        所  西安市雁塔区朱雀大街南段 18 号",
        "法定代表人  张建国",
        "注册资本  壹亿元整",
        "经营范围  电力工程施工、输变电设备安装、机电设备销售。",
        "联系电话  029-87547622",
    ]
    for index, line in enumerate(lines):
        draw.text((80, 170 + index * 66), line, font=mid, fill="black")
    license_img.save(OUT / "license.png")
    print("已生成", OUT / "license.png")

    # —— 名片 ——
    card = Image.new("RGB", (960, 540), "white")
    draw = ImageDraw.Draw(card)
    draw.text((60, 50), "西安明珠电力安装工程有限公司", font=big, fill="black")
    draw.text((60, 130), "王海涛  项目经理", font=mid, fill="black")
    for index, line in enumerate(
        ["手机：138 0013 8000", "电话：029-8888 6666", "邮箱：wang@mingzhu.com", "地址：西安市高新区科技路 33 号"]
    ):
        draw.text((60, 210 + index * 62), line, font=mid, fill="black")
    card.save(OUT / "namecard.png")
    print("已生成", OUT / "namecard.png")

    # —— 报价单 ——
    from docx import Document

    document = Document()
    document.add_heading("报价单", level=1)
    for text in [
        "供应商：深圳市凯歌电子有限公司",
        "纳税人识别号：9144030071526726X9",
        "地址：深圳市南山区科技园北区 5 栋",
        "电话：0755-26551234",
        "联系人：李工",
    ]:
        document.add_paragraph(text)
    table = document.add_table(rows=1, cols=6)
    for cell, text in zip(table.rows[0].cells, ["产品名称", "规格型号", "单位", "数量", "单价", "金额"]):
        cell.text = text
    for row in [
        ("电缆", "YJV-3x35", "米", "100", "25.5", "2550"),
        ("配电箱", "XL-21", "台", "2", "1800", "3600"),
        ("桥架", "200x100", "米", "50", "68", "3400"),
    ]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    document.save(OUT / "quote.docx")
    print("已生成", OUT / "quote.docx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
