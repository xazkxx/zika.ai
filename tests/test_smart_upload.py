"""多文件智能识别（smart-upload）的接口级回归（2026-10-23）。

## 守的四件事

1. **只读不写库**：识别只是"帮着填表"，绝不能在库里留下任何档案；
2. **Word 表格能出明细**：报价单的商品行是靠 `python-docx` 的表格解析，不是靠正则猜；
3. **多文件合并只补空缺**：第二份文件里更含糊的值不许盖掉第一份认到的；
4. **认不出就说认不出**：不支持的格式、空文件、坏了的中文名，都要有可读的 error，
   而不是 500 或"静默返回空字段"。

OCR（图片）单独一条用例：模型首次加载要十几秒，且依赖本机装了 rapidocr_onnxruntime。
没装就 skip，不让它把整个测试集拖红。
"""

import io
import uuid

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.services.smart_upload import (
    parse_items_from_text,
    parse_party_fields,
    parse_product_fields,
)

client = TestClient(app)
BASE = "/api/v1/smart-upload"


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        rows = len(table)
        cols = len(table[0])
        table_obj = document.add_table(rows=rows, cols=cols)
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                table_obj.rows[r].cells[c].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


QUOTE_PARAGRAPHS = [
    "报价单",
    "供应商：深圳市凯歌电子有限公司",
    "纳税人识别号：9144030071526726X9",
    "地址：深圳市南山区科技园北区 5 栋",
    "电话：0755-26551234",
    "联系人：李工",
]
QUOTE_TABLE = [
    ["产品名称", "规格型号", "单位", "数量", "单价", "金额"],
    ["电缆", "YJV-3x35", "米", "100", "25.5", "2550"],
    ["配电箱", "XL-21", "台", "2", "1800", "3600"],
]


def _post(files: list[tuple[str, bytes]], kind: str = "customer", mode: str = "parse"):
    payload = [("files", (name, content, "application/octet-stream")) for name, content in files]
    return client.post(f"{BASE}/parse", files=payload, data={"kind": kind, "mode": mode})


def test_meta_lists_formats_and_fields():
    response = client.get(f"{BASE}/meta")
    assert response.status_code == 200, response.text
    body = response.json()
    assert set(body["kinds"]) == {"customer", "supplier", "product"}
    for suffix in (".png", ".jpg", ".pdf", ".docx"):
        assert suffix in body["accept"]
    assert ".exe" not in body["accept"]
    assert "tax_no" in body["fields"]["supplier"]
    assert body["max_files"] >= 3


def test_docx_quote_extracts_fields_and_items():
    payload = _docx(QUOTE_PARAGRAPHS, QUOTE_TABLE)
    response = _post([("报价单.docx", payload)], kind="supplier")
    assert response.status_code == 200, response.text
    body = response.json()

    assert body["files"][0]["ok"] is True
    assert body["files"][0]["engine"] == "python-docx"
    assert body["fields"]["name"] == "深圳市凯歌电子有限公司", body["fields"]
    assert body["fields"]["tax_no"] == "9144030071526726X9"
    # 表格表头「产品名称」不能被当成公司名
    assert "规格型号" not in body["fields"]["name"]

    assert len(body["items"]) == 2, body["items"]
    first = body["items"][0]
    assert first["name"] == "电缆"
    assert first["unit"] == "米"
    assert first["quantity"] == 100.0
    assert first["unit_price"] == 25.5
    assert body["field_labels"]["name"] == "名称"


def test_parse_never_writes_any_record():
    """识别只是填表的辅助：调完接口，客户/供应商一条都不许多。"""
    before = client.get("/api/v1/customers?limit=500").json()
    before_rows = len(before if isinstance(before, list) else before.get("items") or [])
    payload = _docx(QUOTE_PARAGRAPHS, QUOTE_TABLE)
    _post([("报价单.docx", payload)], kind="customer")
    after = client.get("/api/v1/customers?limit=500").json()
    after_rows = len(after if isinstance(after, list) else after.get("items") or [])
    assert after_rows == before_rows


def test_multi_file_merge_only_fills_gaps():
    """两份文件：第一份有完整地址，第二份只有"见附件" —— 不能被覆盖。"""
    first = _docx(["供应商：深圳市凯歌电子有限公司", "地址：深圳市南山区科技园北区 5 栋"])
    second = _docx(["地址：见附件", "电话：0755-26551234"])
    body = _post([("a.docx", first), ("b.docx", second)], kind="supplier").json()
    assert body["fields"]["name"] == "深圳市凯歌电子有限公司"
    assert body["fields"]["address"] == "深圳市南山区科技园北区 5 栋"
    assert body["fields"]["phone"] == "0755-26551234"


def test_attach_mode_skips_recognition():
    payload = _docx(QUOTE_PARAGRAPHS, QUOTE_TABLE)
    body = _post([("报价单.docx", payload)], kind="customer", mode="attach").json()
    assert body["fields"] == {}
    assert body["items"] == []
    assert body["files"][0]["ok"] is True  # 至少要能读出内容，说明文件是好的
    assert "保存" in body["message"]


def test_unsupported_suffix_is_reported_not_crashed():
    body = _post([("病毒.exe", b"MZ\x90\x00")], kind="customer").json()
    assert body["files"][0]["ok"] is False
    assert "不支持" in (body["files"][0]["error"] or "")
    assert body["fields"] == {}


def test_empty_file_is_reported():
    body = _post([("空文件.docx", b"")], kind="customer").json()
    assert body["files"][0]["ok"] is False
    assert "空" in (body["files"][0]["error"] or "")


def test_bad_kind_and_mode_are_400():
    payload = _docx(["客户：测试"])
    assert _post([("a.docx", payload)], kind="robot").status_code == 400
    assert _post([("a.docx", payload)], kind="customer", mode="magic").status_code == 400


def test_too_many_files_is_400():
    payload = _docx(["客户：测试"])
    files = [(f"f{index}.docx", payload) for index in range(7)]
    assert _post(files, kind="customer").status_code == 400


def test_tax_number_does_not_fake_a_mobile():
    """18 位统一社会信用代码里能凑出 `1[3-9]\\d{9}` —— 早先会被认成手机号。"""
    fields, _evidence, _warnings = parse_party_fields(
        "深圳市凯歌电子有限公司\n纳税人识别号：9144030071526726X9\n电话：0755-26551234",
        "supplier",
    )
    assert fields.get("mobile") is None
    assert fields["phone"] == "0755-26551234"


def test_name_falls_back_to_filename():
    """执照图片糊了、文字没读出来时，文件名「…有限公司营业执照.jpg」是最后的线索。"""
    body = _post(
        [("陕西送变电工程有限公司营业执照.png", b"not-an-image")],
        kind="customer",
    ).json()
    # 图片是坏的 → 读取失败，但文件名里的公司名仍然被采纳
    assert body["files"][0]["ok"] is False
    assert body["fields"].get("name") == "陕西送变电工程有限公司"


def test_text_item_lines_are_parsed():
    rows = parse_items_from_text(
        "1  电缆  YJV-3x35  米  100  25.5  2550\n"
        "2  配电箱  XL-21  台  2  1800  3600\n"
    )
    assert [row["name"] for row in rows] == ["电缆", "配电箱"]
    assert rows[0]["quantity"] == 100.0


@pytest.mark.skipif(
    not __import__("importlib").util.find_spec("rapidocr_onnxruntime"),
    reason="本机没装 rapidocr_onnxruntime，跳过真实 OCR 用例",
)
def test_real_ocr_on_license_image():
    """真跑一次 OCR：把执照文字画在图上，识别结果必须认得出公司名与税号。

    这条用例慢（首次加载模型十几秒），但它是唯一能证明"图片这条路真的通"的用例 ——
    mock 掉 OCR 等于什么都没验。
    """
    import os

    from PIL import Image, ImageDraw, ImageFont

    font_path = "C:/Windows/Fonts/msyh.ttc"
    if not os.path.exists(font_path):
        pytest.skip("没有中文字体，画不出可识别的执照图")

    font = ImageFont.truetype(font_path, 30)
    image = Image.new("RGB", (1000, 620), "white")
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(
        [
            "营 业 执 照",
            "统一社会信用代码  91610103MA6TX8K23Q",
            "名        称  陕西送变电工程有限公司",
            "住        所  西安市雁塔区朱雀大街南段 18 号",
            "法定代表人  张建国",
            "联系电话  029-87547622",
        ]
    ):
        draw.text((60, 40 + index * 80), line, font=font, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")

    body = _post([("执照.png", buffer.getvalue())], kind="customer").json()
    assert body["files"][0]["ok"] is True, body["files"][0]
    assert body["fields"].get("tax_no") == "91610103MA6TX8K23Q", body["fields"]
    name = body["fields"].get("name", "")
    assert "陕西送变电工程有限公司" in name, name
    assert body["fields"].get("contact") == "张建国"


# ---------------------------------------------------------------------------
# 商品：字段解析 + 用明细首行兜底（2026-10-23 铺开第二步）
# ---------------------------------------------------------------------------
def test_product_fields_from_text():
    fields, evidence, warnings = parse_product_fields(
        "产品资料\n"
        "产品名称：低压电缆\n"
        "规格型号：YJV-3x35\n"
        "单位：米\n"
        "销售价：¥25.50\n"
        "采购价：21.80\n"
        "增值税税率：13%\n"
        "生产厂家：西安明珠线缆有限公司\n"
    )
    assert fields["name"] == "低压电缆"
    assert fields["standards"] == "YJV-3x35"
    assert fields["unit"] == "米"
    assert fields["sale_price"] == "25.50"   # ¥ 与千分位都要洗掉
    assert fields["purchase_price"] == "21.80"
    assert fields["tax_rate"] == "13"
    assert fields["manufacturer"] == "西安明珠线缆有限公司"
    assert evidence["name"]
    assert warnings == []


def test_product_fields_warn_when_missing():
    """只有一行"资料标题"的文档：不该把标题当商品名，且要给出补录提醒。"""
    fields, _evidence, warnings = parse_product_fields("产品资料\n")
    assert "name" not in fields, fields
    assert any("产品名称" in warning for warning in warnings)


def test_product_kind_api_uses_first_table_row_as_fallback():
    """报价单常常是"抬头没写产品名，明细全在表里" —— 用第一行兜底。"""
    payload = _docx(["产品报价单（2026 版）"], QUOTE_TABLE)
    body = _post([("报价单.docx", payload)], kind="product").json()
    assert body["files"][0]["ok"] is True
    fields = body["fields"]
    assert fields["name"] == "电缆", fields
    assert fields["standards"] == "YJV-3x35", fields
    assert fields["unit"] == "米", fields
    assert fields["sale_price"] == "25.5", fields
    # 明细行仍然单独返回（前端可以按行选）
    assert len(body["items"]) == 2


def test_product_kind_api_prefers_explicit_fields_over_table():
    payload = _docx(
        ["产品名称：配电箱\n规格型号：XL-21\n单位：台\n销售价：￥1800\n"],
        QUOTE_TABLE,
    )
    fields = _post([("产品.docx", payload)], kind="product").json()["fields"]
    assert fields["name"] == "配电箱"
    assert fields["standards"] == "XL-21"
    assert fields["sale_price"] == "1800"
