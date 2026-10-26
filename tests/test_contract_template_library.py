"""合同模板库 + 完整合同打包 的接口级回归（2026-10-22）。

## 守的四件事

1. **上传→列表→下载**闭环；后缀不认识的文件必须被拒（不能先落盘再报错）；
2. `contracts.template_id` 真的落库，且选了不存在的模板要被拦（否则以后静默退化）；
3. **完整合同**按模板类型分派：Word 走变量替换并返回 .docx、PDF 合并页数增加、
   图片追加一页；
4. **模板删了，历史合同照样能下载正文**（法律文件不能被模板清理连累）。
"""

import io
import uuid
from datetime import date

import pymupdf as fitz
from docx import Document
from PIL import Image
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)
TODAY = date.today().isoformat()


def _uniq(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:8]}"


def _customer() -> dict:
    response = client.post(
        "/api/v1/customers",
        json={"code": _uniq("CUST"), "name": f"测试模板客户{uuid.uuid4().hex[:6]}"},
    )
    assert response.status_code == 201, response.text
    return response.json()


def _product() -> dict:
    response = client.post(
        "/api/v1/products", json={"sku": _uniq("SKU"), "name": "测试模板商品", "unit": "件"}
    )
    assert response.status_code == 201, response.text
    return response.json()


def _contract(customer: dict, product: dict, template_id: int | None = None) -> dict:
    body = {
        "direction": "out",
        "party_id": customer["id"],
        "party_name": customer["name"],
        "sign_date": TODAY,
        "delivery_date": "2026-11-01",
        "delivery_place": "西安工地A区",
        "items": [{"product_id": product["id"], "quantity": 10, "unit_price": 99.5}],
    }
    if template_id:
        body["template_id"] = template_id
    response = client.post("/api/v1/contracts", json=body)
    assert response.status_code == 201, response.text
    return response.json()


def _docx_template(suffix: str) -> bytes:
    doc = Document()
    doc.add_paragraph(f"甲方：{{{{party_name}}}}（{suffix}）")
    doc.add_paragraph("交货地点：{{delivery_place}}  日期：{{delivery_date}}")
    doc.add_paragraph("总价：{{total_amount}}（{{total_chinese}}）")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _upload(name: str, filename: str, payload: bytes, note: str = "") -> dict:
    response = client.post(
        "/api/v1/contract-templates",
        files={"file": (filename, payload, "application/octet-stream")},
        data={"name": name, "note": note},
    )
    assert response.status_code == 201, response.text
    return response.json()


def _pdf_template(pages: int = 2) -> bytes:
    doc = fitz.open()
    for index in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"TEMPLATE PAGE {index + 1}")
    raw = doc.tobytes()
    doc.close()
    return raw


def _png() -> bytes:
    image = Image.new("RGB", (600, 800), "white")
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


# ===========================================================================
# 一、模板上传 / 列表 / 下载 / 删除
# ===========================================================================
def test_upload_list_and_flags():
    created = _upload("标准工矿产品购销合同", "标准工矿.docx", _docx_template("v1"), note="2026 版")
    assert created["kind"] == "word"
    assert created["supports_variables"] is True
    assert created["previewable"] is False
    assert created["file_size"] > 0

    rows = client.get("/api/v1/contract-templates").json()
    assert any(row["id"] == created["id"] for row in rows)

    # 原件能读回来，且内容一致
    file_resp = client.get(f"/api/v1/contract-templates/{created['id']}/file")
    assert file_resp.status_code == 200
    assert file_resp.content[:2] == b"PK"  # docx 是 zip 包

    assert client.delete(f"/api/v1/contract-templates/{created['id']}").status_code == 204
    assert client.get(f"/api/v1/contract-templates/{created['id']}/file").status_code == 404


def test_previewable_flags_for_pdf_and_image():
    """PDF / 图片能在界面里预览（预览按钮只对它们出现），Word 不能在线预览。"""
    pdf = _upload("PDF 模板", "p.pdf", _pdf_template(1))
    image = _upload("图片模板", "p.png", _png())
    try:
        assert pdf["kind"] == "pdf" and pdf["previewable"] is True
        assert image["kind"] == "image" and image["previewable"] is True
        assert pdf["supports_variables"] is False
    finally:
        client.delete(f"/api/v1/contract-templates/{pdf['id']}")
        client.delete(f"/api/v1/contract-templates/{image['id']}")


def test_reject_unknown_extension():
    response = client.post(
        "/api/v1/contract-templates",
        files={"file": ("bad.exe", b"MZ....", "application/octet-stream")},
        data={"name": "坏文件"},
    )
    assert response.status_code == 400, response.text
    assert "不支持" in response.text


# ===========================================================================
# 二、template_id 落库与校验
# ===========================================================================
def test_contract_keeps_template_id_and_rejects_missing_template():
    customer = _customer()
    product = _product()
    template = _upload("模板A", "a.docx", _docx_template("A"))

    contract = _contract(customer, product, template["id"])
    assert int(contract["template_id"]) == int(template["id"])

    bad = client.post(
        "/api/v1/contracts",
        json={
            "direction": "out",
            "party_id": customer["id"],
            "party_name": customer["name"],
            "sign_date": TODAY,
            "template_id": 99999999,
            "items": [{"product_id": product["id"], "quantity": 1, "unit_price": 1}],
        },
    )
    assert bad.status_code == 404, bad.text
    assert "模板" in bad.text

    client.delete(f"/api/v1/contracts/{contract['id']}?force=true")
    client.delete(f"/api/v1/contract-templates/{template['id']}")


# ===========================================================================
# 三、完整合同：三种模板三条路径
# ===========================================================================
def test_full_contract_word_replaces_variables():
    customer = _customer()
    product = _product()
    template = _upload("Word 模板", "word.docx", _docx_template("W"))
    contract = _contract(customer, product, template["id"])

    response = client.get(f"/api/v1/contracts/{contract['id']}/full-contract")
    assert response.status_code == 200, response.text
    assert "wordprocessingml" in response.headers["content-type"]

    content = Document(io.BytesIO(response.content))
    text = "\n".join(p.text for p in content.paragraphs)
    assert customer["name"] in text, "客户名没替换进去"
    assert "西安工地A区" in text, "交货地点没替换进去"
    assert "995.00" in text, "总额没替换进去"
    assert "{{" not in text, "还有没替换掉的占位符"
    assert int(response.headers.get("X-Template-Vars-Replaced", "0")) >= 2

    client.delete(f"/api/v1/contracts/{contract['id']}?force=true")
    client.delete(f"/api/v1/contract-templates/{template['id']}")


def test_full_contract_pdf_merges_pages():
    customer = _customer()
    product = _product()
    template = _upload("PDF 模板", "pdf.pdf", _pdf_template(2))
    contract = _contract(customer, product, template["id"])

    response = client.get(f"/api/v1/contracts/{contract['id']}/full-contract")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"

    merged = fitz.open(stream=response.content, filetype="pdf")
    try:
        # 正文至少 1 页 + 模板 2 页
        assert merged.page_count >= 3, f"合并后只有 {merged.page_count} 页"
    finally:
        merged.close()

    client.delete(f"/api/v1/contracts/{contract['id']}?force=true")
    client.delete(f"/api/v1/contract-templates/{template['id']}")


def test_full_contract_image_appends_one_page():
    customer = _customer()
    product = _product()
    template = _upload("图片模板", "scan.png", _png())
    contract = _contract(customer, product, template["id"])

    response = client.get(f"/api/v1/contracts/{contract['id']}/full-contract")
    assert response.status_code == 200
    doc = fitz.open(stream=response.content, filetype="pdf")
    try:
        assert doc.page_count >= 2
    finally:
        doc.close()

    client.delete(f"/api/v1/contracts/{contract['id']}?force=true")
    client.delete(f"/api/v1/contract-templates/{template['id']}")


def test_full_contract_without_template_is_plain_pdf():
    customer = _customer()
    product = _product()
    contract = _contract(customer, product, None)

    response = client.get(f"/api/v1/contracts/{contract['id']}/full-contract")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content[:4] == b"%PDF"

    client.delete(f"/api/v1/contracts/{contract['id']}?force=true")


def test_deleted_template_does_not_break_history():
    """模板删了，历史合同照样能下载正文 —— 合同是法律文件，不能被模板清理连累。"""
    customer = _customer()
    product = _product()
    template = _upload("临时模板", "tmp.docx", _docx_template("T"))
    contract = _contract(customer, product, template["id"])

    assert client.delete(f"/api/v1/contract-templates/{template['id']}").status_code == 204

    response = client.get(f"/api/v1/contracts/{contract['id']}/full-contract")
    assert response.status_code == 200, response.text
    # 模板没了 → 退化成系统范本 PDF（而不是 500）
    assert response.headers["content-type"] == "application/pdf"

    client.delete(f"/api/v1/contracts/{contract['id']}?force=true")


# ===========================================================================
# 四、无变量兜底 + 幽灵记录清扫（2026-10-24）
# ===========================================================================
def test_full_contract_word_without_variables_prepends_summary():
    """模板里一个 {{变量}} 都没有时，完整合同必须**先**给一页『表单数据汇总页』，
    模板原件从第二页开始 —— 用户拿到手的不能是个空壳原件。"""
    doc = Document()
    doc.add_paragraph("窗帘购销合同（纸质范本，没有占位符）")
    doc.add_paragraph("第一条：标的为窗帘及配件。")
    buffer = io.BytesIO()
    doc.save(buffer)

    customer = _customer()
    product = _product()
    template = _upload("无变量 Word", "plain.docx", buffer.getvalue())
    contract = _contract(customer, product, template["id"])

    response = client.get(f"/api/v1/contracts/{contract['id']}/full-contract")
    assert response.status_code == 200, response.text
    assert "wordprocessingml" in response.headers["content-type"]
    assert response.headers.get("X-Template-Vars-Replaced") == "0"

    content = Document(io.BytesIO(response.content))
    texts = [p.text for p in content.paragraphs if p.text.strip()]
    # 汇总页在最前
    assert texts[0].startswith("表单数据汇总页"), f"第一段是：{texts[0]!r}"
    assert customer["name"] in texts[1], "汇总页没列客户"
    # 模板原件内容在汇总页之后
    joined = "\n".join(texts)
    assert "窗帘购销合同" in joined, "模板原件丢了"
    assert texts.index("窗帘购销合同（纸质范本，没有占位符）") > 0
    summary_end = max(i for i, t in enumerate(texts) if "标的物明细" in t or "测试模板商品" in t)
    template_start = texts.index("窗帘购销合同（纸质范本，没有占位符）")
    assert template_start > summary_end, "模板原件没有排在汇总页后面"

    client.delete(f"/api/v1/contracts/{contract['id']}?force=true")
    client.delete(f"/api/v1/contract-templates/{template['id']}")


def test_ghost_template_swept_from_list():
    """磁盘文件被手工删掉的"幽灵记录"，列表接口要自动清掉，
    且 file/variables 接口返回统一的中文防呆提示。"""
    from app.core import file_storage

    template = _upload("幽灵模板", "ghost.docx", _docx_template("G"))
    # file_path 是内部字段（_serialize 不外泄），直接查库拿
    from app.db.session import SessionLocal
    from app.models.contract_template import ContractTemplate as _CT

    _db = SessionLocal()
    try:
        _row = _db.get(_CT, template["id"])
        file_path = _row.file_path
    finally:
        _db.close()
    assert file_storage.delete_file(file_path) is True

    # file / variables 接口：记录还在、文件没了 → 统一中文防呆提示
    file_resp = client.get(f"/api/v1/contract-templates/{template['id']}/file")
    assert file_resp.status_code == 404
    assert file_resp.json()["detail"] == "该模板文件已损坏或丢失，请重新上传"

    var_resp = client.get(f"/api/v1/contract-templates/{template['id']}/variables")
    assert var_resp.status_code == 404
    assert var_resp.json()["detail"] == "该模板文件已损坏或丢失，请重新上传"

    # 列表：幽灵记录被自动清扫，不再出现
    rows = client.get("/api/v1/contract-templates").json()
    assert all(r["id"] != template["id"] for r in rows), "幽灵记录还在列表里"

    # 清理 DB 残留（列表清扫已删，这里兜底幂等）
    client.delete(f"/api/v1/contract-templates/{template['id']}")
