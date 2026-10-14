"""Excel 导入导出测试：模板下载 / 导入（新建·更新·跳过·预检·部分失败）/ 导出。

测试数据命名遵循 scripts/clean_test_data.py 的识别签名（CUST-XLSX-xxxx / SKU-XLSX-xxxx
/ SUP-xxxx / WH-xxxx），跑完由 conftest 自动清理，不污染业务库。
"""

import io
import uuid

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.services import excel_io

client = TestClient(app)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _hex8() -> str:
    return uuid.uuid4().hex[:8]


def build_upload(kind: str, records: list[dict]) -> bytes:
    """以标准模板为底，删掉示例行后写入数据，返回可上传的 xlsx 字节。"""
    tmpl = client.get(f"/api/v1/excel/templates/{kind}")
    assert tmpl.status_code == 200
    wb = load_workbook(io.BytesIO(tmpl.content))
    ws = wb["导入数据"]
    ws.delete_rows(3)  # 第 3 行是示例行
    specs = excel_io.COLUMNS[kind]
    for rec in records:
        ws.append([rec.get(col["field"]) for col in specs])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(kind: str, records: list[dict], *, mode: str = "upsert", dry_run: bool = False):
    return client.post(
        f"/api/v1/excel/import/{kind}",
        params={"mode": mode, "dry_run": str(dry_run).lower()},
        files={"file": ("import.xlsx", build_upload(kind, records), XLSX_MIME)},
    )


# ---------------------------------------------------------------------------
# 模板与列表
# ---------------------------------------------------------------------------

def test_template_download_all_kinds():
    for kind in ("products", "customers", "suppliers", "warehouses", "opening-stock"):
        resp = client.get(f"/api/v1/excel/templates/{kind}")
        assert resp.status_code == 200, kind
        assert resp.headers["content-type"].startswith(XLSX_MIME)
        assert resp.content[:2] == b"PK"  # xlsx 是 zip 包
        wb = load_workbook(io.BytesIO(resp.content))
        assert "导入数据" in wb.sheetnames
        assert "填写说明" in wb.sheetnames
        ws = wb["导入数据"]
        headers = [c.value for c in ws[2]]
        assert any("编码" in str(h) for h in headers)


def test_template_unknown_kind_404():
    assert client.get("/api/v1/excel/templates/no-such").status_code == 404


def test_kinds_list():
    resp = client.get("/api/v1/excel/kinds")
    assert resp.status_code == 200
    kinds = {item["kind"] for item in resp.json()}
    assert {"products", "customers", "suppliers", "warehouses", "opening-stock"} <= kinds
    customers = next(i for i in resp.json() if i["kind"] == "customers")
    assert "客户编码" in customers["required_fields"]


# ---------------------------------------------------------------------------
# 导入：档案
# ---------------------------------------------------------------------------

def test_import_creates_customers():
    tag = _hex8()
    recs = [
        {"code": f"CUST-XLSX-{tag}", "name": "导入测试客户甲", "phone": "029-0001", "level": "A"},
        {"code": f"CUST-XLSX-{_hex8()}", "name": "导入测试客户乙", "credit_limit": 5000},
    ]
    resp = upload("customers", recs)
    assert resp.status_code == 200
    body = resp.json()
    assert body["created"] == 2 and body["failed"] == 0, body
    assert body["dry_run"] is False

    listed = client.get("/api/v1/customers").json()
    codes = {row["code"] for row in listed}
    assert recs[0]["code"] in codes
    # 字段确实落库了
    row = next(r for r in listed if r["code"] == recs[0]["code"])
    assert row["name"] == "导入测试客户甲"
    assert row["phone"] == "029-0001"


def test_import_update_then_skip():
    tag = _hex8()
    code = f"CUST-XLSX-{tag}"

    first = upload("customers", [{"code": code, "name": "导入更新客户-初始"}]).json()
    assert first["created"] == 1

    # upsert：同编码 → 更新
    second = upload("customers", [{"code": code, "name": "导入更新客户-改名"}]).json()
    assert second["updated"] == 1 and second["created"] == 0, second
    listed = client.get("/api/v1/customers").json()
    assert next(r for r in listed if r["code"] == code)["name"] == "导入更新客户-改名"

    # skip：同编码 → 跳过，名称不变
    third = upload("customers", [{"code": code, "name": "不该生效的名字"}], mode="skip").json()
    assert third["skipped"] == 1 and third["updated"] == 0, third
    listed = client.get("/api/v1/customers").json()
    assert next(r for r in listed if r["code"] == code)["name"] == "导入更新客户-改名"


def test_import_dry_run_does_not_persist():
    code = f"CUST-XLSX-{_hex8()}"
    resp = upload("customers", [{"code": code, "name": "预检客户不该落库"}], dry_run=True)
    assert resp.status_code == 200
    body = resp.json()
    assert body["dry_run"] is True and body["created"] == 1

    listed = client.get("/api/v1/customers").json()
    assert code not in {row["code"] for row in listed}


def test_import_partial_failure_keeps_good_rows():
    """一行缺必填不应拖垮整批——这是批量导入最容易踩的坑。"""
    tag1, tag2 = _hex8(), _hex8()
    recs = [
        {"code": f"CUST-XLSX-{tag1}", "name": "部分失败-好行一"},
        {"code": "", "name": "部分失败-缺编码"},          # 必填为空
        {"code": f"CUST-XLSX-{tag2}", "name": "部分失败-好行二"},
    ]
    body = upload("customers", recs).json()
    assert body["created"] == 2, body
    assert body["failed"] == 1, body
    assert body["errors"] and "必填" in body["errors"][0]["message"]

    codes = {row["code"] for row in client.get("/api/v1/customers").json()}
    assert f"CUST-XLSX-{tag1}" in codes and f"CUST-XLSX-{tag2}" in codes


def test_import_bad_mode_400():
    resp = client.post(
        "/api/v1/excel/import/customers",
        params={"mode": "explode"},
        files={"file": ("x.xlsx", build_upload("customers", []), XLSX_MIME)},
    )
    assert resp.status_code == 400


def test_import_not_xlsx_400():
    resp = client.post(
        "/api/v1/excel/import/customers",
        files={"file": ("x.xlsx", b"this is not a spreadsheet", XLSX_MIME)},
    )
    assert resp.status_code == 400


# ---------------------------------------------------------------------------
# 导入：期初库存（联动库存）
# ---------------------------------------------------------------------------

def test_import_opening_stock_links_inventory():
    tag = _hex8()
    wh_code, sku = f"WH-{tag}", f"SKU-XLSX-{tag}"
    wh = client.post("/api/v1/warehouses", json={"code": wh_code, "name": f"主仓库XLSX{tag}"})
    assert wh.status_code == 201, wh.text
    wh_id = wh.json()["id"]
    pr = client.post("/api/v1/products", json={"sku": sku, "name": "期初导入测试商品", "unit": "件"})
    assert pr.status_code == 201, pr.text
    prod_id = pr.json()["id"]

    # 首次导入 → 新建
    rec = {"warehouse_code": wh_code, "sku": sku, "quantity": 100, "unit_price": 3.5}
    body = upload("opening-stock", [rec]).json()
    assert body["created"] == 1 and body["failed"] == 0, body

    stock = client.get("/api/v1/inventory/stock").json()
    rows = stock["rows"] if isinstance(stock, dict) and "rows" in stock else stock
    row = next(r for r in rows if r["product_id"] == prod_id and r["warehouse_id"] == wh_id)
    assert row["quantity"] == 100

    # 再次导入同键 → 覆盖更新，不重复累加
    body2 = upload("opening-stock", [dict(rec, quantity=160)]).json()
    assert body2["updated"] == 1 and body2["created"] == 0, body2
    stock = client.get("/api/v1/inventory/stock").json()
    rows = stock["rows"] if isinstance(stock, dict) and "rows" in stock else stock
    row = next(r for r in rows if r["product_id"] == prod_id and r["warehouse_id"] == wh_id)
    assert row["quantity"] == 160

    # mode=skip → 保持不变
    body3 = upload("opening-stock", [dict(rec, quantity=999)], mode="skip").json()
    assert body3["skipped"] == 1, body3
    stock = client.get("/api/v1/inventory/stock").json()
    rows = stock["rows"] if isinstance(stock, dict) and "rows" in stock else stock
    row = next(r for r in rows if r["product_id"] == prod_id and r["warehouse_id"] == wh_id)
    assert row["quantity"] == 160


def test_import_opening_stock_unknown_refs_reported():
    tag = _hex8()
    body = upload(
        "opening-stock",
        [{"warehouse_code": f"WH-{tag}", "sku": f"SKU-XLSX-{tag}", "quantity": 5}],
    ).json()
    assert body["failed"] == 1 and body["created"] == 0, body
    assert "不存在" in body["errors"][0]["message"]


# ---------------------------------------------------------------------------
# 导出
# ---------------------------------------------------------------------------

def test_export_contains_created_rows():
    code = f"CUST-XLSX-{_hex8()}"
    upload("customers", [{"code": code, "name": "导出测试客户"}])

    resp = client.get("/api/v1/excel/export/customers")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(XLSX_MIME)
    assert "attachment" in resp.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "客户编码" in headers
    idx = headers.index("客户编码")
    codes = {ws.cell(row=r, column=idx + 1).value for r in range(2, ws.max_row + 1)}
    assert code in codes


def test_export_unknown_kind_404():
    assert client.get("/api/v1/excel/export/no-such").status_code == 404
